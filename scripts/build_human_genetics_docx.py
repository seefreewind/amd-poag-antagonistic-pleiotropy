#!/usr/bin/env python3
"""Build a Word submission draft from the audited Human Genetics Markdown package."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
MANUSCRIPT = ROOT / "manuscript/HUMAN_GENETICS_MANUSCRIPT_V1.md"
TABLES = ROOT / "manuscript/HUMAN_GENETICS_TABLES.md"
LEGENDS = ROOT / "manuscript/HUMAN_GENETICS_FIGURE_LEGENDS.md"
OUT = ROOT / "manuscript/HUMAN_GENETICS_MANUSCRIPT_V1.docx"

FIGURES = {
    1: ROOT / "results/figures/Figure1_global_rg.tiff",
    2: ROOT / "results/figures/Figure2_local_rg.tiff",
    3: ROOT / "results/figures/Figure3_LOCO.tiff",
    4: ROOT / "results/figures/Figure4_phase4_local_rg_comparison.png",
    5: ROOT / "results/figures/Figure5_L006_L007_direction_resolved_adjudication.png",
    6: ROOT / "results/figures/Figure6_phase4_architecture_summary.png",
}

BLACK = "000000"
BLUE = "2F6F8F"
LIGHT_BLUE = "EAF1F5"
LIGHT_GRAY = "F4F6F7"
BORDER = "D9D9D9"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = BORDER, size: str = "6") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep(paragraph, keep_next=False, keep_lines=True) -> None:
    paragraph.paragraph_format.keep_with_next = keep_next
    paragraph.paragraph_format.keep_together = keep_lines


def remove_paragraph_borders(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def remove_style_borders(style) -> None:
    p_pr = style._element.pPr
    if p_pr is not None:
        border = p_pr.find(qn("w:pBdr"))
        if border is not None:
            p_pr.remove(border)


def add_run_markup(paragraph, text: str, size=11, bold=False, italic=False, color=BLACK) -> None:
    """Small inline Markdown subset for the audited manuscript."""
    pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            run.font.size = Pt(size)
            run.font.name = "Times New Roman"
            run.font.color.rgb = RGBColor.from_string(color)
        token = match.group(0)
        if token.startswith("**"):
            content, b, i = token[2:-2], True, False
        elif token.startswith("*"):
            content, b, i = token[1:-1], False, True
        else:
            content, b, i = token[1:-1], False, False
        run = paragraph.add_run(content)
        run.bold = bold or b
        run.italic = italic or i
        run.font.size = Pt(size)
        run.font.name = "Times New Roman"
        run.font.color.rgb = RGBColor.from_string(color)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        run.bold = bold
        run.italic = italic
        run.font.size = Pt(size)
        run.font.name = "Times New Roman"
        run.font.color.rgb = RGBColor.from_string(color)


def add_body_paragraph(doc: Document, text: str, style="Body Text", size=11, space_after=6):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.08
    add_run_markup(p, text, size=size)
    return p


def parse_table_section(name: str):
    text = TABLES.read_text()
    m = re.search(rf"^## {re.escape(name)}$", text, re.M)
    if not m:
        raise ValueError(f"Missing table section: {name}")
    n = re.search(r"^## ", text[m.end():], re.M)
    block = text[m.end():m.end() + n.start() if n else len(text)]
    lines = [line.strip() for line in block.splitlines() if line.strip()]
    table_lines = [line for line in lines if line.startswith("|")]
    rows = []
    for line in table_lines:
        cells = [x.strip() for x in line.strip("|").split("|")]
        if all(set(x) <= {"-", ":", " "} for x in cells):
            continue
        rows.append(cells)
    notes = [line for line in lines if not line.startswith("|")]
    return rows, notes


def add_table(doc: Document, number: int, title: str, section_name: str) -> None:
    rows, notes = parse_table_section(section_name)
    cap = doc.add_paragraph(style="Caption")
    cap.paragraph_format.space_before = Pt(8)
    cap.paragraph_format.space_after = Pt(4)
    run = cap.add_run(f"Table {number} {title}")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)
    set_keep(cap, keep_next=True)

    tbl = doc.add_table(rows=1, cols=len(rows[0]))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = True
    header = tbl.rows[0]
    set_repeat_table_header(header)
    for idx, value in enumerate(rows[0]):
        cell = header.cells[idx]
        cell.text = ""
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, BLUE)
        set_cell_borders(cell)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(value)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(7.3)
        r.font.color.rgb = RGBColor(255, 255, 255)
    for ridx, row_values in enumerate(rows[1:]):
        row = tbl.add_row()
        for cidx, value in enumerate(row_values):
            cell = row.cells[cidx]
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, LIGHT_GRAY if ridx % 2 else "FFFFFF")
            set_cell_borders(cell)
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if cidx in (0, 1, 7, 11) else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(value)
            r.font.name = "Arial"
            r.font.size = Pt(7.1 if number != 3 else 6.9)
            r.font.color.rgb = RGBColor.from_string(BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    for note in notes:
        p = doc.add_paragraph(style="Caption")
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(7)
        add_run_markup(p, note, size=8)


def parse_legends():
    text = LEGENDS.read_text()
    out = {}
    for match in re.finditer(r"\*\*Fig\. (\d)\*\* (.*?)(?=\n\n\*\*Fig\.|\Z)", text, re.S):
        out[int(match.group(1))] = re.sub(r"\s+", " ", match.group(2).strip())
    return out


def add_figure(doc: Document, number: int, legends: dict[int, str]) -> None:
    path = FIGURES[number]
    if not path.exists():
        raise FileNotFoundError(path)
    # Keep the final architecture-summary figure and its relatively long legend
    # together.  Without this page break Word can leave the last line of the
    # caption stranded at the top of the following page.
    if number == 6:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(6.25))
    cap = doc.add_paragraph(style="Caption")
    cap.paragraph_format.keep_together = True
    cap.paragraph_format.space_after = Pt(8)
    add_run_markup(cap, legends[number], size=8.5)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(6)
    body = styles["Body Text"]
    body.font.name = "Times New Roman"
    body.font.size = Pt(11)
    body.font.color.rgb = RGBColor.from_string(BLACK)
    body.paragraph_format.first_line_indent = Inches(0.2)
    body.paragraph_format.line_spacing = 1.08
    body.paragraph_format.space_after = Pt(6)
    title = styles["Title"]
    title.font.name = "Times New Roman"
    title.font.size = Pt(16)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_after = Pt(12)
    remove_style_borders(title)
    for name, size, before, after in (("Heading 1", 13, 14, 6), ("Heading 2", 12, 10, 4), ("Heading 3", 11, 8, 3)):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    caption = styles["Caption"]
    caption.font.name = "Times New Roman"
    caption.font.size = Pt(8.5)
    caption.font.color.rgb = RGBColor.from_string(BLACK)
    caption.paragraph_format.line_spacing = 1.0


def add_heading(doc: Document, level: int, text: str):
    p = doc.add_paragraph(style=f"Heading {min(level, 3)}")
    add_run_markup(p, text, size={1: 13, 2: 12, 3: 11}[min(level, 3)], bold=True)
    set_keep(p, keep_next=True)
    return p


def build() -> None:
    doc = Document()
    configure_styles(doc)
    sec = doc.sections[0]
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.8)
    sec.left_margin = Inches(0.85)
    sec.right_margin = Inches(0.85)
    sec.header_distance = Inches(0.35)
    sec.footer_distance = Inches(0.35)
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = header.add_run("Human Genetics Research Article")
    hr.font.name = "Arial"
    hr.font.size = Pt(8)
    hr.font.color.rgb = RGBColor.from_string("666666")
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("AMD–POAG direction-resolved genetic analysis")
    fr.font.name = "Arial"
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor.from_string("666666")

    legends = parse_legends()
    lines = MANUSCRIPT.read_text().splitlines()
    i = 0
    paragraph_lines: list[str] = []
    inserted = set()

    def flush():
        nonlocal paragraph_lines
        if not paragraph_lines:
            return
        text = " ".join(x.strip() for x in paragraph_lines).strip()
        paragraph_lines = []
        if not text:
            return
        if text.startswith("**Keywords:**"):
            p = doc.add_paragraph(style="Body Text")
            p.paragraph_format.first_line_indent = Inches(0)
            add_run_markup(p, text, size=10)
        elif text.startswith("**Authors:**"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(6)
            set_keep(p, keep_next=True)
            add_run_markup(p, text, size=11)
        elif text.startswith("**Affiliation 1:**") or text.startswith("**Affiliation 2:**"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(3)
            set_keep(p, keep_next=True)
            add_run_markup(p, text, size=9.5)
        elif text.startswith("**Corresponding author:**"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(12)
            add_run_markup(p, text, size=9.5)
        elif text.startswith("[AUTHOR INPUT") or text.startswith("[AUTHOR CONFIRMATION") or text.startswith("[CODE AVAILABILITY"):
            p = add_body_paragraph(doc, text, style="Body Text", size=10)
            p.paragraph_format.first_line_indent = Inches(0)
        else:
            add_body_paragraph(doc, text)

    def after_marker(marker: str):
        if marker in inserted:
            return
        if marker == "table1":
            add_table(doc, 1, "GWAS datasets and analytical roles", "Table 1 | GWAS datasets and analytical roles")
        elif marker == "table2":
            add_table(doc, 2, "Global genetic correlation and robustness analyses", "Table 2 | Global genetic correlation and robustness analyses")
        elif marker == "table3":
            add_table(doc, 3, "Shared loci, direction, fine-mapping support and final evidence category", "Table 3 | Shared loci, direction, fine-mapping support and final evidence category")
        elif marker.startswith("fig"):
            add_figure(doc, int(marker[3:]), legends)
        inserted.add(marker)

    while i < len(lines):
        line = lines[i]
        if line.startswith("# "):
            flush()
            p = doc.add_paragraph(style="Title")
            remove_paragraph_borders(p)
            set_keep(p, keep_next=True)
            add_run_markup(p, line[2:].strip(), size=16, bold=True)
            i += 1
            continue
        if line.startswith("## "):
            flush()
            heading = line[3:].strip()
            add_heading(doc, 1, heading)
            i += 1
            continue
        if line.startswith("### "):
            flush()
            add_heading(doc, 2, line[4:].strip())
            i += 1
            continue
        if not line.strip():
            flush()
            i += 1
            continue
        if line.startswith("|---") or line.startswith("|---"):
            i += 1
            continue
        paragraph_lines.append(line)
        joined = " ".join(paragraph_lines)
        if joined.startswith("The analytical dataset comprised"):
            flush(); after_marker("table1")
        elif joined.startswith("In the primary no-MHC LDSC analysis") and "Fig. 1" in joined:
            flush(); after_marker("fig1")
        elif joined.startswith("The primary estimate also provides the scale"):
            pass
        elif joined.startswith("HDL-L evaluated 2,463") and "Fig. 3" in joined:
            flush(); after_marker("table2"); after_marker("fig2"); after_marker("fig3")
        elif joined.startswith("The frozen cross-trait discovery stage") and "Table 3" in joined:
            flush(); after_marker("table3")
        elif joined.startswith("The unified coloc.susie analysis") and "Fig. 5" in joined:
            flush(); after_marker("fig4"); after_marker("fig5")
        elif joined.startswith("The final synthesis is presented in Fig. 6"):
            flush(); after_marker("fig6")
        i += 1
    flush()

    # Add a document property note for the author review stage without inventing metadata.
    props = doc.core_properties
    props.title = "Direction-resolved genetic analysis of age-related macular degeneration and primary open-angle glaucoma"
    props.subject = "Human Genetics Research Article submission draft"
    props.author = "[AUTHOR INPUT NEEDED]"
    props.comments = "Generated from the audited Human Genetics Markdown manuscript; author declarations remain placeholders."
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
